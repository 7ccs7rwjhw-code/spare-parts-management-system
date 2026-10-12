import io, os, tempfile, threading, hashlib
from pathlib import Path
from decimal import Decimal
from unittest import skipUnless
from django.test import TestCase, TransactionTestCase, Client, override_settings
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError, PermissionDenied
from django.db import connection, close_old_connections
from openpyxl import Workbook, load_workbook
from .models import *
from .excel import *
from .services import *

TEMPLATE = Path(__file__).resolve().parents[2] / "templates/default21.xlsx"
SOURCE = Path(
    os.getenv(
        "SOURCE_DIR",
        str(Path(__file__).resolve().parent.parent.parent / "task_package/source"),
    )
)
ARCHIVE = Path(
    os.getenv(
        "ARCHIVE_DIR", str(Path(__file__).resolve().parent.parent.parent / "archive")
    )
)


def sample(
    width=21,
    qty="200m",
    title="备件购买",
    month=10,
    code="000001",
    description="示例物料",
):
    w = Workbook()
    s = w.active
    s.title = title
    if width == 21:
        headers = list(HEADERS)
        headers[11] = "9月份实际"
        headers[12] = f"{month}月预估"
    elif width == 18:
        headers = [HEADERS[i] for i in [1, 2, 3, 4, 5, 6, 7, 8, 9, 10]] + [
            f"{month}月预估",
            "未交",
            "库存",
            "未交",
            "库存",
            "客户别",
            "科长",
            "",
        ]
    elif width == 17:
        headers = [HEADERS[i] for i in range(11)] + [
            f"{month}月预估",
            "未交",
            "库存",
            "客户别",
            "科长",
            "全厂库存",
        ]
    else:
        headers = ["机种", "品类", "数量", "单价", "总价"]
    s.append(headers)
    values = [""] * width
    if width in [21, 17, 18]:
        kind, m = adapter(headers)
        values[m["new_material_code"]] = code
        values[m["material_description"]] = description
        values[m["forecast_qty"]] = qty
        for k in ["customer", "section_manager"]:
            values[m[k]] = k
        if width == 18:
            values[11:15] = ["11", "12", "13", "14"]
            values[17] = "unlabelled"
    else:
        values = ["M", "P", "4", "5", "20"]
    s.append(values)
    b = io.BytesIO()
    w.save(b)
    return b.getvalue()


class Base:
    def init_users(self):
        U = get_user_model()
        self.admin = U.objects.create_user(
            "admin-test", password="TestPass!739", is_superuser=True, is_staff=True
        )
        self.owner = U.objects.create_user("owner", password="TestPass!739")
        Profile.objects.create(user=self.owner, department="A")
        self.other = U.objects.create_user("other", password="TestPass!739")
        Profile.objects.create(user=self.other, department="B")
        self.reviewer = U.objects.create_user("reviewer", password="TestPass!739")
        Profile.objects.create(user=self.reviewer, role="reviewer", department="A")
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.over = override_settings(MEDIA_ROOT=self.tmp.name)
        self.over.enable()
        self.addCleanup(self.over.disable)

    def batch(self):
        return Batch.objects.create(
            year=2026,
            month=10,
            kind="spare_purchase",
            requester=self.owner,
            department="A",
        )

    def line(self, b, qty="20", unit="个"):
        vals = {k: "" for k in FIELDS}
        vals.update(
            material_description="示例物料",
            new_material_code="000001",
            forecast_qty=qty,
        )
        return save_line(self.owner, b.pk, b.lock_version, vals, unit, confirmed=True)

    def approve(self, b):
        b = transition(self.owner, b.pk, b.lock_version, "submit")
        b = transition(self.reviewer, b.pk, b.lock_version, "approve")
        select_effective(self.admin, b.pk, b.lock_version, reason="测试选择")
        b.refresh_from_db()
        return b


class ImportTests(Base, TestCase):
    def setUp(self):
        self.init_users()

    def test_IMP01_21_forecast_column(self):
        r, s = scan(sample(), "202610月.xlsx")
        self.assertEqual(r[0]["mapped"]["forecast_qty"], "200m")
        self.assertEqual(r[0]["template"], "21")

    def test_IMP02_department_category(self):
        f, _ = ingest(sample(title="部门自购"), "sample.xlsx", self.admin)
        b = confirm_sheet(
            self.admin, f.pk, "部门自购", 2026, 10, self.owner, "A", "人工确认"
        )
        self.assertEqual(b.kind, "department_purchase")

    def test_IMP03_18_requires_mapping(self):
        r, _ = scan(sample(18), "8月.xlsx")
        self.assertEqual(r[0]["mapped"]["forecast_qty"], "200m")
        self.assertNotIn("factory_stock", r[0]["mapped"])
        f, _ = ingest(sample(18), "8月.xlsx", self.admin)
        with self.assertRaises(ValidationError):
            confirm_sheet(
                self.admin, f.pk, "备件购买", 2026, 8, self.owner, "A", "月份确认"
            )
        b = confirm_sheet(
            self.admin,
            f.pk,
            "备件购买",
            2026,
            8,
            self.owner,
            "A",
            "业务确认M/L",
            {"factory_stock": 12, "outstanding_qty": 11},
        )
        self.assertEqual(b.lines.get().factory_stock, Decimal("12"))
        self.assertEqual(b.lines.get().outstanding_qty, Decimal("11"))

    def test_IMP04_17_headers(self):
        r, _ = scan(sample(17), "9月.xlsx")
        m = r[0]["mapped"]
        self.assertEqual(m["forecast_qty"], "200m")
        self.assertEqual(m["customer"], "customer")
        self.assertEqual(m["section_manager"], "section_manager")

    def test_IMP05_other_template(self):
        f, _ = ingest(sample(5), "tools.xlsx", self.admin)
        self.assertEqual(f.rows.get().template, "other")
        with self.assertRaises(ValidationError):
            confirm_sheet(
                self.admin, f.pk, "备件购买", 2026, 10, self.owner, "A", "理由"
            )
        self.assertEqual(Batch.objects.count(), 0)

    def test_IMP06_quantity_raw_and_unit(self):
        f, _ = ingest(sample(), "sample.xlsx", self.admin)
        b = confirm_sheet(
            self.admin, f.pk, "备件购买", 2026, 10, self.owner, "A", "依据"
        )
        l = b.lines.get()
        self.assertEqual(l.raw_quantity, "200m")
        self.assertEqual(l.quantity, Decimal("200"))
        self.assertEqual(l.unit, "m")
        self.assertFalse(l.confirmed)

    def test_IMP07_no_code_retained(self):
        f, _ = ingest(sample(code="无料号下单"), "sample.xlsx", self.admin)
        b = confirm_sheet(
            self.admin, f.pk, "备件购买", 2026, 10, self.owner, "A", "依据"
        )
        self.assertIsNone(b.lines.get().item)
        self.assertEqual(b.lines.get().values["new_material_code"], "无料号下单")

    def test_IMP08_hash_idempotent(self):
        data = sample()
        f, _ = ingest(data, "one.xlsx", self.admin)
        b = confirm_sheet(
            self.admin, f.pk, "备件购买", 2026, 10, self.owner, "A", "依据"
        )
        f2, new = ingest(data, "one.xlsx", self.admin)
        self.assertFalse(new)
        self.assertEqual(f.pk, f2.pk)
        self.assertEqual(Batch.objects.count(), 1)
        self.assertEqual(Line.objects.count(), 1)

    def test_IMP09_different_versions_not_auto_effective(self):
        for qty in ["20", "30"]:
            f, _ = ingest(sample(qty=qty), "same.xlsx", self.admin)
            confirm_sheet(
                self.admin, f.pk, "备件购买", 2026, 10, self.owner, "A", "依据"
            )
        self.assertEqual(SourceFile.objects.count(), 2)
        self.assertFalse(Batch.objects.filter(effective=True).exists())
        self.assertEqual(list(summary(self.admin, 2026, 10)), [])

    def test_IMP09_same_name_shared_version_family(self):
        batches = []
        for qty in ["20", "30"]:
            f, _ = ingest(sample(qty=qty), "same.xlsx", self.admin)
            b = confirm_sheet(
                self.admin, f.pk, "备件购买", 2026, 10, self.owner, "A", "版本确认"
            )
            l = b.lines.get()
            b = save_line(
                self.owner, b.pk, b.lock_version, l.values, "个", l.pk, confirmed=True
            )
            b = self.approve(b)
            batches.append(b)
        self.assertEqual(batches[0].family, batches[1].family)
        self.assertEqual(Batch.objects.filter(effective=True).count(), 1)
        self.assertEqual(list(summary(self.admin, 2026, 10))[0]["total"], Decimal(30))

    def test_IMP10_paths_preserved(self):
        data = sample()
        f, _ = ingest(data, "8月/file.xlsx", self.admin)
        f2, new = ingest(data, "9月/file.xlsx", self.admin)
        self.assertEqual(f.pk, f2.pk)
        self.assertEqual(f.locations.count(), 2)

    def test_IMP11_month_conflicts_require_confirmation(self):
        r, _ = scan(sample(month=7), "26年/6月/20268月份.xlsx")
        self.assertTrue(any("冲突" in x for x in r[0]["issues"]))
        f, _ = ingest(sample(month=7), "20268月份.xlsx", self.admin)
        self.assertFalse(Batch.objects.exists())
        with self.assertRaises(ValidationError):
            confirm_sheet(self.admin, f.pk, "备件购买", 2026, 8, self.owner, "A", "")

    def test_IMP12_invalid_actual_month(self):
        w = load_workbook(io.BytesIO(sample()))
        w.active["L1"] = "0月份实际"
        b = io.BytesIO()
        w.save(b)
        r, _ = scan(b.getvalue(), "sample.xlsx")
        self.assertTrue(any("非法月份" in x for x in r[0]["issues"]))

    def test_IMP13_quantity_states(self):
        for raw, reason in [("", "数量缺失"), ("-", "数量占位"), ("bad", "非法数量")]:
            d, u, e = number(raw)
            self.assertIsNone(d)
            self.assertIn(reason, e)
        self.assertEqual(number("0")[0], Decimal(0))

    def test_IMP14_excel_error(self):
        f, _ = ingest(sample(qty="#N/A"), "err.xlsx", self.admin)
        self.assertTrue(f.rows.get().issues)
        b = confirm_sheet(
            self.admin, f.pk, "备件购买", 2026, 10, self.owner, "A", "错误保留"
        )
        self.assertIsNone(b.lines.get().quantity)

    def test_IMP15_unmatched_retained(self):
        f, _ = ingest(sample(code="unmatched"), "sample.xlsx", self.admin)
        b = confirm_sheet(
            self.admin, f.pk, "备件购买", 2026, 10, self.owner, "A", "依据"
        )
        self.assertIsNone(b.lines.get().item)
        self.assertEqual(
            b.lines.get().source_row.source.sha256,
            hashlib.sha256(sample(code="unmatched")).hexdigest(),
        )

    def test_IMP16_partial_rows_reported(self):
        w = load_workbook(io.BytesIO(sample(qty="20")))
        w.active.append(["", "X", "Invalid"] + [""] * 9 + ["bad"])
        buf = io.BytesIO()
        w.save(buf)
        f, _ = ingest(buf.getvalue(), "partial.xlsx", self.admin)
        self.assertEqual(f.rows.count(), 2)
        self.client.force_login(self.admin)
        r = self.client.get(f"/imports/{f.pk}/csv/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("非法", r.content.decode())

    def test_upload_security(self):
        with self.assertRaises(ValueError):
            checked_workbook(b"not xlsx")
        with self.assertRaises(ValueError):
            checked_workbook(b"x" * (25 * 1024 * 1024 + 1))

    @skipUnless((SOURCE / "设备备件清单-z.xlsx").exists(), "Provide private SOURCE_DIR")
    def test_DATA01_to_DATA06_real_catalog(self):
        path = SOURCE / "设备备件清单-z.xlsx"
        data = path.read_bytes()
        h = hashlib.sha256(data).hexdigest()
        r, s = scan(data, "catalog.xlsx", True)
        self.assertEqual(len(s), 15)
        f, _ = ingest(data, "catalog.xlsx", self.admin, True)
        self.assertEqual(Item.objects.count(), 0)
        apply_catalog(f.pk, self.admin)
        self.assertEqual(Item.objects.count(), 292)
        self.assertEqual(Item.objects.filter(new_code="").count(), 120)
        from django.db.models import Count

        groups = (
            Item.objects.exclude(new_code="")
            .values("new_code")
            .annotate(n=Count("id"))
            .filter(n__gt=1)
        )
        self.assertEqual(groups.count(), 7)
        self.assertTrue(all(x["n"] == 2 for x in groups))
        self.assertEqual(f.rows.filter(template="usage").count(), 182)
        self.assertEqual(Batch.objects.count(), 0)
        self.assertTrue(
            Alias.objects.filter(kind="candidate", source_row__sheet="ASM").exists()
        )
        self.assertTrue(Item.objects.filter(new_code="", old_code__gt="").exists())
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), h)

    @skipUnless(ARCHIVE.exists(), "Provide private ARCHIVE_DIR")
    def test_real_archive_inventory(self):
        files = list(ARCHIVE.rglob("*.xlsx"))
        self.assertEqual(len(files), 14)
        candidates = 0
        for p in sorted(files):
            f, new = ingest(
                p.read_bytes(), p.relative_to(ARCHIVE).as_posix(), self.admin
            )
            candidates += (
                f.rows.filter(template__in=["21", "17", "18"])
                .exclude(status="non_business_row")
                .count()
            )
        self.assertEqual(SourceFile.objects.count(), 13)
        self.assertEqual(candidates, 195)
        self.assertEqual(
            SourceRow.objects.filter(template__in=["21", "17", "18"])
            .exclude(status="non_business_row")
            .count(),
            194,
        )
        self.assertTrue(
            SourceRow.objects.filter(
                mapped__forecast_qty="200m", mapped__parsed_unit="m"
            ).exists()
        )
        self.assertEqual(SourceLocation.objects.count(), 14)

    def test_formula_cached_evidence_and_unchanged_confirmation(self):
        data = sample(qty="2")
        w = load_workbook(io.BytesIO(data))
        w.active["K2"] = "=1+1"
        buf = io.BytesIO()
        w.save(buf)
        f, _ = ingest(buf.getvalue(), "formula.xlsx", self.admin)
        row = f.rows.get()
        self.assertEqual(row.original_values[10], "=1+1")
        self.assertEqual(row.cached_values[10], "")
        b = confirm_sheet(
            self.admin,
            f.pk,
            "备件购买",
            2026,
            10,
            self.owner,
            "A",
            "公式无缓存保留待确认",
        )
        l = b.lines.get()
        b = save_line(
            self.owner, b.pk, b.lock_version, l.values, "个", l.pk, confirmed=True
        )
        self.assertIsNone(b.lines.get().theoretical_monthly_qty)
        self.assertEqual(b.lines.get().values["theoretical_monthly_qty"], "=1+1")


class WorkflowTests(Base, TestCase):
    def setUp(self):
        self.init_users()

    def test_FLOW01_draft_restored(self):
        b = self.line(self.batch(), "0.125")
        l = b.lines.get()
        self.assertEqual(l.quantity, Decimal("0.125"))
        self.client.force_login(self.owner)
        self.assertEqual(self.client.get(f"/batches/{b.pk}/").status_code, 200)

    def test_FLOW02_copy_clears_snapshots(self):
        b = self.line(self.batch())
        l = b.lines.get()
        l.values.update(
            factory_stock="10", outstanding_qty="30", previous_month_actual_qty="1"
        )
        l.factory_stock = 10
        l.save()
        n = clone(self.owner, b.pk, b.lock_version, 2026, 11)
        self.assertEqual(n.month, 11)
        self.assertEqual(n.status, "draft")
        self.assertIsNone(n.lines.get().factory_stock)
        self.assertEqual(n.lines.get().values["factory_stock"], "")
        self.assertEqual(b.lines.get().values["factory_stock"], "10")

    def test_FLOW03_submit_locks(self):
        b = self.line(self.batch())
        b = transition(self.owner, b.pk, b.lock_version, "submit")
        with self.assertRaises(ValidationError):
            self.line(b)
        self.assertEqual(b.status, "submitted")

    def test_FLOW04_return_reason(self):
        b = self.line(self.batch())
        b = transition(self.owner, b.pk, b.lock_version, "submit")
        with self.assertRaises(ValidationError):
            transition(self.reviewer, b.pk, b.lock_version, "return", "")
        b = transition(self.reviewer, b.pk, b.lock_version, "return", "请核对数量")
        self.assertEqual(b.status, "returned")
        self.assertEqual(b.events.last().comment, "请核对数量")
        self.assertEqual(b.events.last().actor, self.reviewer)

    def test_FLOW05_approved_revision(self):
        b = self.approve(self.line(self.batch()))
        n = clone(self.owner, b.pk, b.lock_version, 2026, 10, True)
        self.assertEqual(n.parent, b)
        self.assertEqual(n.family, b.family)
        self.assertEqual(n.version, 2)
        self.assertTrue(b.effective)
        self.assertFalse(n.effective)
        with self.assertRaises(ValidationError):
            self.line(b)
        self.assertTrue(n.events.get().before["lines"])

    def test_FLOW06_object_permissions(self):
        b = self.line(self.batch())
        old = b.lines.get().quantity
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(f"/batches/{b.pk}/").status_code, 404)
        self.assertEqual(
            self.client.post(
                f"/batches/{b.pk}/action/submit/", {"lock_version": b.lock_version}
            ).status_code,
            403,
        )
        self.assertEqual(b.lines.get().quantity, old)

    def test_FLOW07_admin_endpoint_permission(self):
        self.client.force_login(self.owner)
        self.assertEqual(self.client.get("/imports/").status_code, 403)
        self.assertTrue(Audit.objects.filter(action="denied_admin").exists())

    def test_FLOW08_stale_version(self):
        b = self.batch()
        n = self.line(b)
        with self.assertRaises(Conflict):
            self.line(b, "999")
        self.assertEqual(n.lines.count(), 1)

    def test_FLOW09_versions_only_selected(self):
        b = self.approve(self.line(self.batch()))
        n = clone(self.owner, b.pk, b.lock_version, 2026, 10, True)
        l = n.lines.get()
        n = save_line(
            self.owner, n.pk, n.lock_version, l.values, l.unit, l.pk, confirmed=True
        )
        n = self.approve(n)
        b.refresh_from_db()
        self.assertFalse(b.effective)
        self.assertEqual(list(summary(self.admin, 2026, 10))[0]["total"], Decimal("20"))

    def test_FLOW10_units_not_summed(self):
        b = self.line(self.batch(), "200m", "m")
        b = self.line(b, "20", "个")
        self.approve(b)
        totals = list(summary(self.admin, 2026, 10))
        self.assertEqual(len(totals), 2)
        self.assertEqual(
            {r["unit"]: r["total"] for r in totals},
            {"m": Decimal(200), "个": Decimal(20)},
        )

    def test_reviewer_department_and_self_approval(self):
        b = self.line(self.batch())
        b = transition(self.owner, b.pk, b.lock_version, "submit")
        with self.assertRaises(PermissionDenied):
            transition(self.other, b.pk, b.lock_version, "approve")
        self.assertFalse(visible(self.other).filter(pk=b.pk).exists())

    def test_csrf_and_login(self):
        c = Client(enforce_csrf_checks=True)
        self.assertEqual(c.get("/").status_code, 302)
        c.force_login(self.owner)
        self.assertEqual(
            c.post(
                "/batches/new/", {"year": 2026, "month": 10, "kind": "spare_purchase"}
            ).status_code,
            403,
        )

    def test_all_pages_render(self):
        self.client.force_login(self.admin)
        b = self.batch()
        f, _ = ingest(sample(), "sample.xlsx", self.admin)
        for path in [
            "/",
            "/items/",
            "/items/new/",
            "/imports/",
            f"/imports/{f.pk}/",
            "/reports/",
            "/queue/",
            "/audit/",
            f"/batches/{b.pk}/",
            f"/batches/{b.pk}/line/",
            f"/batches/{b.pk}/copy/",
        ]:
            self.assertEqual(self.client.get(path).status_code, 200, path)


class ExportTests(Base, TestCase):
    def setUp(self):
        self.init_users()

    def workbook(self, b, month=10):
        return load_workbook(
            io.BytesIO(
                export_workbook(Batch.objects.filter(pk=b.pk), 2026, month, TEMPLATE)
            )
        )

    def test_EXPORT01_two_sheets_21(self):
        b = self.line(self.batch())
        w = self.workbook(b)
        self.assertEqual(w.sheetnames, ["备件购买", "部门自购"])
        self.assertEqual(w["备件购买"].max_column, 21)
        self.assertEqual(w["备件购买"]["M1"].value, "10月预估")

    def test_EXPORT02_dynamic_month(self):
        b = self.line(self.batch())
        w = self.workbook(b, 1)
        self.assertEqual(w["备件购买"]["L1"].value, "12月份实际")
        self.assertEqual(w["备件购买"]["M1"].value, "1月预估")

    def test_EXPORT03_all_fields_roundtrip(self):
        b = self.line(self.batch(), "200m", "m")
        b = self.line(b, "20", "个")
        l = b.lines.first()
        v = dict(l.values)
        v.update(old_material_code="0012", remark="原备注", order_number="00004")
        b = save_line(self.owner, b.pk, b.lock_version, v, l.unit, l.pk, confirmed=True)
        w = self.workbook(b)
        s = w["备件购买"]
        self.assertEqual(s.max_row, 3)
        self.assertEqual(s["A2"].value, "0012")
        self.assertEqual(s["T2"].value, "原备注")
        self.assertEqual(number(s["M2"].value)[:2], (Decimal("200"), "m"))
        self.assertTrue(
            all(
                s.cell(1, j).value == h
                for j, h in enumerate(HEADERS, 1)
                if j not in [12, 13]
            )
        )
        self.assertEqual(w["部门自购"]["C2"].value, None)

    def test_EXPORT04_formula_injection(self):
        b = self.line(self.batch())
        l = b.lines.get()
        v = dict(l.values)
        v["material_description"] = "=1+1"
        v["remark"] = "@test"
        b = save_line(self.owner, b.pk, b.lock_version, v, l.unit, l.pk, confirmed=True)
        s = self.workbook(b)["备件购买"]
        self.assertEqual(s["C2"].value, "=1+1")
        self.assertEqual(s["C2"].data_type, "s")
        self.assertEqual(s["T2"].data_type, "s")

    def test_EXPORT05_leading_zeros(self):
        b = self.line(self.batch())
        c = self.workbook(b)["备件购买"]["B2"]
        self.assertEqual(c.value, "000001")
        self.assertEqual(c.data_type, "s")
        self.assertEqual(c.number_format, "@")

    def test_template_no_original_content(self):
        w = load_workbook(TEMPLATE)
        self.assertEqual(w.sheetnames, ["备件购买", "部门自购"])
        for s in w:
            self.assertTrue(
                all(c.value is None for row in s.iter_rows(min_row=2) for c in row)
            )
            self.assertFalse(s._images)
            self.assertFalse(s._charts)


class ConcurrentTests(Base, TransactionTestCase):
    reset_sequences = True

    def setUp(self):
        self.init_users()

    def test_FLOW08_real_postgresql_parallel_save(self):
        self.assertEqual(connection.vendor, "postgresql")
        b = self.batch()
        gate = threading.Barrier(2)
        results = []

        def writer(q):
            close_old_connections()
            try:
                u = get_user_model().objects.get(pk=self.owner.pk)
                gate.wait(timeout=10)
                vals = {k: "" for k in FIELDS}
                vals.update(material_description="测试", forecast_qty=str(q))
                save_line(u, b.pk, 1, vals, "个", confirmed=True)
                results.append("saved")
            except Conflict:
                results.append("conflict")
            except Exception as e:
                results.append(type(e).__name__ + str(e))
            finally:
                close_old_connections()

        ts = [threading.Thread(target=writer, args=(q,)) for q in [10, 20]]
        for t in ts:
            t.start()
        for t in ts:
            t.join(timeout=15)
        self.assertCountEqual(results, ["saved", "conflict"])
        self.assertEqual(b.lines.count(), 1)

    def test_IMP08_parallel_identical_upload(self):
        gate = threading.Barrier(2)
        results = []
        data = sample()

        def writer():
            close_old_connections()
            try:
                u = get_user_model().objects.get(pk=self.admin.pk)
                gate.wait(timeout=10)
                f, new = ingest(data, "same.xlsx", u)
                results.append(new)
            except Exception as e:
                results.append(str(e))
            finally:
                close_old_connections()

        ts = [threading.Thread(target=writer) for _ in range(2)]
        for t in ts:
            t.start()
        for t in ts:
            t.join(timeout=15)
        self.assertCountEqual(results, [True, False])
        self.assertEqual(SourceFile.objects.count(), 1)

    def test_DEPLOY05_five_users_parallel_independent(self):
        users = [
            get_user_model().objects.create_user("parallel-" + str(i)) for i in range(5)
        ]
        batches = [
            Batch.objects.create(
                year=2026, month=10, kind="spare_purchase", requester=u, department="A"
            )
            for u in users
        ]
        gate = threading.Barrier(5)
        results = []

        def writer(i):
            close_old_connections()
            try:
                u = get_user_model().objects.get(pk=users[i].pk)
                gate.wait(timeout=10)
                vals = {k: "" for k in FIELDS}
                vals.update(
                    material_description="parallel-" + str(i), forecast_qty=str(i + 1)
                )
                save_line(u, batches[i].pk, 1, vals, "个", confirmed=True)
                results.append(visible(u).get().pk == batches[i].pk)
            except Exception as e:
                results.append(str(e))
            finally:
                close_old_connections()

        ts = [threading.Thread(target=writer, args=(i,)) for i in range(5)]
        for t in ts:
            t.start()
        for t in ts:
            t.join(timeout=15)
        self.assertEqual(results, [True] * 5)
        for i, b in enumerate(batches):
            self.assertEqual(b.lines.get().quantity, Decimal(i + 1))
