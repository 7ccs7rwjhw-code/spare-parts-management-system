"""Pure Excel parsing: no database writes or implicit month corrections."""

import io
import re
import zipfile
from pathlib import Path
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

FIELDS = "old_material_code new_material_code material_description equipment_name material_attribute safety_stock machine_count_per_line usage_per_machine lifetime_months line_name theoretical_monthly_qty previous_month_actual_qty forecast_qty customer section_manager factory_stock outstanding_qty nvb2 required_purchase_qty remark order_number".split()
HEADERS = [
    "原物料",
    "新建料号",
    "物料描述",
    "使用设备",
    "属性",
    "安全库存",
    "设备数量/1条大线",
    "备件使用量/1台设备",
    "使用寿命(以月为单位)",
    "线体",
    "月需求量/理论",
    "上月实际",
    "本月预估",
    "客户别",
    "科长",
    "全厂库存",
    "未交",
    "NVB2",
    "需要请购",
    "备注",
    "单号",
]
NUMERIC = set(FIELDS[5:9] + FIELDS[10:13] + FIELDS[15:19])
INVALID_CODES = {"-", "无料号下单", "无料号", "0"}


def txt(v):
    return "" if v is None else str(v)


def normal(v):
    return re.sub("\\s+", "", txt(v))


def number(raw):
    t = txt(raw).strip()
    if not t:
        return (None, "", "数量缺失")
    if t == "-":
        return (None, "", "数量占位")
    if re.fullmatch("[+-]?(?:\\d+(?:\\.\\d*)?|\\.\\d+)[eE][+-]?\\d+", t):
        return (None, "", "科学计数法数量须人工确认，不能按单位后缀解析")
    m = re.fullmatch("([+]?(?:\\d+(?:\\.\\d*)?|\\.\\d+))\\s*([^\\d\\s].*)?", t)
    if not m:
        return (None, "", "非法数量或公式/Excel错误")
    try:
        d = Decimal(m[1])
        u = m[2] or ""
        if len(u) > 30:
            return (None, "", "单位超出30字符限制")
        if d >= Decimal("1e14") or d.as_tuple().exponent < -6:
            return (None, u, "数量超过精度范围")
        return (d, u, "单位需确认" if u else "")
    except InvalidOperation:
        return (None, "", "非法数量")


def checked_workbook(data):
    if len(data) > 25 * 1024 * 1024:
        raise ValueError("文件超过25MB")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if (
                len(z.infolist()) > 20000
                or sum((x.file_size for x in z.infolist())) > 150 * 1024 * 1024
            ):
                raise ValueError("解压规模超过限制")
            if any(
                (
                    ".." in Path(x.filename.replace("\\", "/")).parts
                    or x.filename.startswith("/")
                    or ":" in x.filename
                    for x in z.infolist()
                )
            ):
                raise ValueError("不安全的文件路径")
        w = load_workbook(io.BytesIO(data), data_only=False)
        if any((s.max_row > 10000 or s.max_column > 100 for s in w)):
            raise ValueError("工作表范围超过限制")
        return w
    except ValueError:
        raise
    except Exception as e:
        raise ValueError("不是有效或可安全读取的xlsx文件") from e


def candidates(path, headers):
    filename = Path(path).name
    y = re.search("(20\\d{2})(1[0-2]|[1-9])月", filename)
    m = y or re.search("(1[0-2]|[1-9])月", filename)
    folders = [int(x) for x in re.findall("(?:^|/)(\\d{1,2})月(?:/|$)", path)]
    h = [int(x) for v in headers for x in re.findall("(\\d{1,2})月预估", normal(v))]
    a = [int(x) for v in headers for x in re.findall("(\\d{1,2})月份?实际", normal(v))]
    return {
        "filename_year": int(y[1]) if y else None,
        "filename_month": int(y[2]) if y else int(m[1]) if m else None,
        "folder_month": folders[-1] if folders else None,
        "header_month": h,
        "actual_month": a,
    }


def adapter(headers):
    n = [normal(x) for x in headers]
    if not any(("物料描述" in x for x in n)) or not any(("月预估" in x for x in n)):
        return ("other", {})
    width = 21 if "NVB2" in n else 18 if n.count("未交") > 1 else 17
    mapping = {}
    for i, h in enumerate(n):
        if re.fullmatch("\\d+月预估", h):
            mapping["forecast_qty"] = i
        elif re.fullmatch("\\d+月份?实际", h):
            mapping["previous_month_actual_qty"] = i
        else:
            for field, title in zip(FIELDS, HEADERS):
                if h and h == normal(title):
                    mapping[field] = i
    if width == 18:
        mapping.pop("outstanding_qty", None)
        mapping.pop("factory_stock", None)
    return (str(width), mapping)


def scan(data, path, catalog=False):
    w = checked_workbook(data)
    cached = load_workbook(io.BytesIO(data), data_only=True)
    out = []
    sheets = []
    for s in w:
        allrows = list(s.iter_rows(values_only=True))
        first = next(
            (i for i, r in enumerate(allrows) if any((v is not None for v in r))), None
        )
        if first is None:
            sheets.append({"sheet": s.title, "template": "empty", "rows": 0})
            continue
        headers = [txt(v) for v in allrows[first]]
        idx = {normal(h): i for i, h in enumerate(headers) if h}
        if catalog:
            kind = (
                "catalog"
                if s.title not in ["Sheet1", "Sheet2", "领用记录"]
                and ("新料号" in idx or "旧料号" in idx)
                else "usage"
                if s.title == "领用记录"
                else "other"
            )
            mapping = {
                k: idx[h]
                for k, h in [
                    ("new_material_code", "新料号"),
                    ("old_material_code", "旧料号"),
                    ("material_description", "厂内物料说明"),
                ]
                if h in idx
            }
        else:
            kind, mapping = adapter(headers)
        dates = candidates(path, headers)
        si = []
        months = [dates["filename_month"], dates["folder_month"]] + dates[
            "header_month"
        ]
        if len(set((x for x in months if x is not None))) > 1:
            si.append("文件名/目录/表头月份冲突，须确认")
        if any(
            (x < 1 or x > 12 for x in dates["actual_month"] + dates["header_month"])
        ):
            si.append("非法月份表头")
        if kind == "18":
            si.append("双库存/未交及无表头列待映射确认")
        start = len(out)
        for ri, rv in enumerate(allrows[first + 1 :], first + 2):
            if not any((v is not None for v in rv)):
                continue
            raw = [txt(v) for v in rv]
            cache = [
                txt(cached[s.title].cell(ri, j).value) for j in range(1, len(rv) + 1)
            ]
            mapped = {k: raw[j] if j < len(raw) else "" for k, j in mapping.items()}
            issues = list(si)
            status = "candidate"
            if kind in ["other", "usage"]:
                status = "other_template" if kind == "other" else "usage_evidence"
            elif not any(
                (
                    mapped.get(k)
                    for k in [
                        "new_material_code",
                        "old_material_code",
                        "material_description",
                    ]
                )
            ):
                status = "non_business_row"
                issues.append("非物料行，完整保留")
            elif kind == "catalog":
                if not mapped.get("new_material_code"):
                    issues.append("无新料号，待确认")
                if s.title == "ASM" and re.match(
                    "5321\\.", mapped.get("material_description", "")
                ):
                    issues.append("ASM描述列包含料号候选，需人工确认")
            else:
                d, u, err = number(mapped.get("forecast_qty"))
                mapped.update(
                    parsed_quantity=str(d) if d is not None else None, parsed_unit=u
                )
                if err:
                    issues.append(err)
                if mapped.get("new_material_code") in INVALID_CODES:
                    issues.append("非正式料号，待关联")
            for key, idx in mapping.items():
                if key in NUMERIC and raw[idx].startswith("="):
                    mapped[key + "_cached"] = cache[idx]
            for j, v in enumerate(rv):
                if isinstance(v, str) and (
                    v.startswith("=") or v in ["#N/A", "#VALUE!", "#DIV/0!", "#REF!"]
                ):
                    issues.append(f"列{j + 1}公式/错误原值保留，未假设为0")
            if issues and status == "candidate":
                status = "needs_confirmation"
            out.append(
                {
                    "sheet": s.title,
                    "row": ri,
                    "original_values": raw,
                    "headers": headers,
                    "cached_values": cache,
                    "mapped": mapped,
                    "issues": issues,
                    "status": status,
                    "template": kind,
                    "month_candidates": dates,
                }
            )
        sheets.append(
            {
                "sheet": s.title,
                "template": kind,
                "rows": len(out) - start,
                "issues": si,
                "candidates": dates,
            }
        )
    if catalog:
        from collections import Counter

        counts = Counter(
            (
                r["mapped"].get("new_material_code")
                for r in out
                if r["template"] == "catalog" and r["mapped"].get("new_material_code")
            )
        )
        for r in out:
            if (
                r["template"] == "catalog"
                and counts.get(r["mapped"].get("new_material_code"), 0) > 1
            ):
                r["issues"].append("重复新料号，保留独立物料/设备关系，合并需人工确认")
                if r["status"] == "candidate":
                    r["status"] = "needs_confirmation"
    return (out, sheets)
