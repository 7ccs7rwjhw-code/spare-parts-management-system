import io
from copy import copy
from openpyxl import load_workbook, Workbook
from apps.inventory.excel import FIELDS, HEADERS


def clean_template(reference, dest):
    """Retain only cell formats, dimensions and printing from a private reference."""
    src = load_workbook(reference)
    out = Workbook()
    out.remove(out.active)
    for title in ["备件购买", "部门自购"]:
        s = src[title]
        d = out.create_sheet(title)
        for row in s.iter_rows(min_row=1, max_row=2, max_col=21):
            for c in row:
                t = d.cell(c.row, c.column)
                for attr in [
                    "font",
                    "fill",
                    "border",
                    "alignment",
                    "protection",
                    "number_format",
                ]:
                    setattr(t, attr, copy(getattr(c, attr)))
        for key, val in s.column_dimensions.items():
            d.column_dimensions[key] = copy(val)
        for key in [1, 2]:
            d.row_dimensions[key] = copy(s.row_dimensions[key])
        d.sheet_format = copy(s.sheet_format)
        d.sheet_properties = copy(s.sheet_properties)
        d.page_setup = copy(s.page_setup)
        d.page_margins = copy(s.page_margins)
        d.print_options = copy(s.print_options)
        d.freeze_panes = "C2"
        d.print_title_rows = "1:1"
        for j, h in enumerate(HEADERS, 1):
            d.cell(1, j, h)
    out.save(dest)


def export_workbook(batches, year, month, template):
    w = load_workbook(template)
    for s in w:
        for j, h in enumerate(HEADERS, 1):
            s.cell(1, j, h)
        s.cell(1, 12, f"{(12 if month == 1 else month - 1)}月份实际")
        s.cell(1, 13, f"{month}月预估")
    counters = {"spare_purchase": 2, "department_purchase": 2}
    for batch in batches.order_by("id"):
        s = w["备件购买" if batch.kind == "spare_purchase" else "部门自购"]
        for line in batch.lines.order_by("id"):
            r = counters[batch.kind]
            counters[batch.kind] += 1
            for j, key in enumerate(FIELDS, 1):
                v = line.values.get(key, "")
                if key == "forecast_qty":
                    v = (
                        str(line.quantity) + (line.unit or "")
                        if line.quantity is not None
                        else line.raw_quantity
                    )
                c = s.cell(r, j)
                c._style = copy(s.cell(2, j)._style)
                if isinstance(v, str):
                    c.value = v
                    c.data_type = "s"
                    c.number_format = "@"
                else:
                    c.value = v
                if key in ["new_material_code", "old_material_code"]:
                    c.number_format = "@"
            s.row_dimensions[r].height = s.row_dimensions[2].height
    for s in w:
        s.auto_filter.ref = f"A1:U{s.max_row}"
        s.print_area = f"A1:U{s.max_row}"
    b = io.BytesIO()
    w.save(b)
    return b.getvalue()
