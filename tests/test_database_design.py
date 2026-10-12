"""New tests for app ownership, database guarantees and candidate integration."""

from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import pytest
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied
from django.db import IntegrityError, transaction
from django.db.models.deletion import ProtectedError
from django.utils import timezone

from apps.accounts.models import Audit, Profile
from apps.approvals.models import Event
from apps.approvals.services import transition
from apps.inventory.importers import ingest, apply_catalog
from apps.inventory.models import (
    Equipment,
    InventorySnapshot,
    Item,
    SourceFile,
    SourceRow,
)
from apps.inventory.services import confirm_sheet
from apps.requests.models import Batch, Line
from apps.requests.services import select_effective
from apps.core.tests import sample


@pytest.fixture
def users(db):
    admin = get_user_model().objects.create_user("schema-admin", is_superuser=True)
    owner = get_user_model().objects.create_user("schema-owner")
    reviewer = get_user_model().objects.create_user("schema-reviewer")
    Profile.objects.create(user=owner, department="A")
    Profile.objects.create(user=reviewer, role="reviewer", department="A")
    return admin, owner, reviewer


@pytest.fixture
def private_media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    return tmp_path


def test_model_ownership():
    for model, app in [
        (Profile, "accounts"),
        (Audit, "accounts"),
        (Equipment, "inventory"),
        (Item, "inventory"),
        (SourceFile, "inventory"),
        (Batch, "requests"),
        (Line, "requests"),
        (Event, "approvals"),
    ]:
        assert model._meta.app_label == app


@pytest.mark.parametrize(
    "invalid",
    [
        {"month": 0},
        {"month": 13},
        {"year": 1999},
        {"year": 2101},
        {"status": "unknown"},
        {"kind": "unknown"},
        {"effective": True},
        {"version": 0},
        {"lock_version": 0},
    ],
)
def test_batch_constraints_enforced_by_postgresql(users, invalid):
    values = {
        "year": 2026,
        "month": 10,
        "kind": "spare_purchase",
        "requester": users[1],
    }
    values.update(invalid)
    with pytest.raises(IntegrityError), transaction.atomic():
        Batch.objects.create(**values)


def test_roles_and_one_profile_per_user(users):
    with pytest.raises(IntegrityError), transaction.atomic():
        Profile.objects.create(user=users[0], role="unknown")
    with pytest.raises(IntegrityError), transaction.atomic():
        Profile.objects.create(user=users[1])


def test_same_material_code_is_not_global_identity(users):
    equipment = Equipment.objects.create(code="EQ-1", name="设备一")
    first = Item.objects.create(new_code="same", equipment_record=equipment)
    second = Item.objects.create(new_code="same")
    assert first.pk != second.pk
    with pytest.raises(ProtectedError):
        equipment.delete()


def test_monthly_snapshot_has_separate_units_and_provenance(users, private_media):
    source, _ = ingest(sample(), "history.xlsx", users[0])
    item = Item.objects.create(description="snapshot")
    values = dict(
        item=item,
        year=2026,
        month=10,
        warehouse="A",
        as_of=timezone.now(),
        source_row=source.rows.get(),
        stock_qty=Decimal("200"),
    )
    snapshot = InventorySnapshot.objects.create(**values, unit="m")
    InventorySnapshot.objects.create(**values, unit="个")
    assert snapshot.source_row.source.sha256 == source.sha256
    with pytest.raises(IntegrityError), transaction.atomic():
        InventorySnapshot.objects.create(**values, unit="m")
    with pytest.raises(IntegrityError), transaction.atomic():
        InventorySnapshot.objects.create(**(values | {"stock_qty": -1}), unit="kg")


def test_line_projection_preserves_all_text_and_decimal(users, private_media):
    source, _ = ingest(
        sample(qty="200m", description="原描述"), "history.xlsx", users[0]
    )
    batch = confirm_sheet(
        users[0], source.pk, "备件购买", 2026, 10, users[1], "A", "已人工确认"
    )
    line = batch.lines.get()
    assert line.material_description == line.values["material_description"] == "原描述"
    assert line.new_material_code == "000001"
    assert line.forecast_qty == Decimal("200") and line.unit == "m"
    assert line.raw_quantity == "200m"
    line.values["remark"] = "业务修正"
    line.save(update_fields=["values"])
    line.refresh_from_db()
    assert line.remark == "业务修正"
    assert line.source_row.original_values[2] == "原描述"
    with pytest.raises(ProtectedError):
        source.delete()


def test_import_authorization_is_in_service_layer(users, private_media):
    with pytest.raises(PermissionDenied):
        ingest(sample(), "forbidden.xlsx", users[1])
    assert not SourceFile.objects.exists()
    assert Audit.objects.filter(action="denied_admin").exists()


def test_failed_preview_rolls_back_database_and_file_storage(users, private_media):
    with patch(
        "apps.inventory.importers.SourceRow.objects.create",
        side_effect=RuntimeError("failure"),
    ):
        with pytest.raises(RuntimeError):
            ingest(sample(), "failed.xlsx", users[0])
    assert not SourceFile.objects.exists()
    assert not list(Path(private_media).rglob("*.xlsx"))


def test_effective_family_merge_allocates_unique_versions(users):
    batches = [
        Batch.objects.create(
            year=2026,
            month=10,
            kind="spare_purchase",
            requester=users[1],
            status="approved",
        )
        for _ in range(2)
    ]
    select_effective(users[0], batches[0].pk, 1, reason="第一个版本")
    select_effective(
        users[0], batches[1].pk, 1, [batches[0].pk], reason="人工确认同一业务单"
    )
    for batch in batches:
        batch.refresh_from_db()
    assert batches[0].family == batches[1].family
    assert batches[0].version != batches[1].version
    assert not batches[0].effective and batches[1].effective
    event = batches[1].events.get(action="select_effective")
    assert {row["version"] for row in event.after["versions"]} == {1, 2}
    assert {row["family"] for row in event.after["versions"]} == {
        str(batches[1].family)
    }
    with pytest.raises(IntegrityError), transaction.atomic():
        Batch.objects.create(
            year=2026,
            month=10,
            kind="spare_purchase",
            requester=users[1],
            family=batches[1].family,
            version=batches[1].version,
        )


def test_workflow_timestamps_and_operation_audit(users):
    batch = Batch.objects.create(
        year=2026, month=10, kind="spare_purchase", requester=users[1], department="A"
    )
    Line.objects.create(
        batch=batch, quantity=1, unit="个", raw_quantity="1", confirmed=True
    )
    batch = transition(users[1], batch.pk, 1, "submit")
    assert batch.submitted_at is not None
    batch = transition(users[2], batch.pk, batch.lock_version, "approve", "确认")
    assert batch.reviewed_at >= batch.submitted_at
    assert batch.events.get(action="approve").comment == "确认"
    assert Audit.objects.filter(target=f"batch:{batch.pk}", action="approve").exists()


def test_catalog_preserves_duplicates_missing_codes_usage_and_asm(users, private_media):
    import io
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NPM+SPG"
    sheet.append(["新料号", "旧料号", "厂内物料说明"])
    sheet.append(["same", "old1", "描述1"])
    sheet.append(["same", "old2", "描述2"])
    sheet.append(["", "old3", "无新码"])
    asm = workbook.create_sheet("ASM")
    asm.append(["新料号", "旧料号", "厂内物料说明", "供应商物料说明"])
    asm.append(["", "old4", "5321.CANDIDATE", "候选描述"])
    usage = workbook.create_sheet("领用记录")
    usage.append(["物料", "数量"])
    usage.append(["历史领用", 3])
    workbook.create_sheet("空表")
    buffer = io.BytesIO()
    workbook.save(buffer)
    source, _ = ingest(buffer.getvalue(), "catalog.xlsx", users[0], catalog=True)
    assert source.rows.count() == 5
    assert apply_catalog(source.pk, users[0]) == 4
    assert apply_catalog(source.pk, users[0]) == 0
    assert Item.objects.filter(new_code="same").count() == 2
    assert Item.objects.filter(new_code="").count() == 2
    assert source.rows.filter(template="usage").count() == 1
    assert source.rows.filter(sheet="ASM").get().original_values[3] == "候选描述"
    assert not Batch.objects.exists()


def test_import_command_previews_without_writes_and_retries_idempotently(
    users, private_media, tmp_path
):
    import io
    import json
    import hashlib
    from django.core.management import call_command

    path = tmp_path / "history.xlsx"
    data = sample()
    path.write_bytes(data)
    original_hash = hashlib.sha256(data).hexdigest()
    output = io.StringIO()
    call_command(
        "import_excel", str(path), actor=users[0].username, preview=True, stdout=output
    )
    preview = json.loads(output.getvalue())
    assert preview["rows"] == 1 and preview["preview_only"]
    assert not SourceFile.objects.exists()
    for source_path in ["8月/history.xlsx", "9月/history.xlsx"]:
        call_command(
            "import_excel",
            str(path),
            actor=users[0].username,
            source_path=source_path,
            stdout=io.StringIO(),
        )
    assert SourceFile.objects.count() == 1
    assert SourceFile.objects.get().locations.count() == 2
    assert not Batch.objects.exists()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == original_hash
